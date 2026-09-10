from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any

import httpx
from docx import Document

from zypg.mcp_tools.mcp_stdio import call_mcp_stdio, default_npx_command
from zypg.storage.files import file_root, relpath

ALLOWED_TYPES = {
    "single_choice",
    "multiple_choice",
    "true_false",
    "fill_in_blank",
    "short_answer",
    "noun_explanation",
    "calculation",
}

_TYPE_ALIASES = {
    "choice": "single_choice",
    "single": "single_choice",
    "mcq": "single_choice",
    "multi": "multiple_choice",
    "tf": "true_false",
    "truefalse": "true_false",
    "fill": "fill_in_blank",
    "blank": "fill_in_blank",
    "qa": "short_answer",
    "essay": "short_answer",
}

_URL_RE = re.compile(r"https?://[^\s)\]>\"']+")


def use_edupaper_mock() -> bool:
    flag = os.environ.get("EDUPAPER_MCP_MOCK", "").strip().lower()
    if flag in {"1", "true", "yes"}:
        return True
    if flag in {"0", "false", "no"}:
        return False
    if os.environ.get("PYTEST_CURRENT_TEST"):
        return True
    from zypg.config import settings

    return bool(settings.edupaper_mcp_mock)


def normalize_exam_questions(questions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for q in questions:
        raw = str(q.get("type") or "short_answer").strip().lower()
        qtype = _TYPE_ALIASES.get(raw, raw)
        if qtype not in ALLOWED_TYPES:
            qtype = "short_answer"
        item: dict[str, Any] = {
            "type": qtype,
            "content": str(q.get("content") or q.get("stem") or ""),
        }
        options = q.get("options")
        if options:
            item["options"] = [str(x) for x in options]
        if q.get("answer") is not None:
            item["answer"] = str(q["answer"])
        points = q.get("points")
        if points is None:
            points = q.get("score")
        if points is not None:
            try:
                item["points"] = float(points)
            except (TypeError, ValueError):
                pass
        out.append(item)
    return out


def extract_download_url(text: str) -> str | None:
    if not text:
        return None
    labeled = re.search(r"Download Link:\s*(\S+)", text, re.I)
    if labeled:
        return labeled.group(1).strip().rstrip("。).,")
    found = _URL_RE.search(text)
    return found.group(0) if found else None


def _safe_stem(title: str) -> str:
    stem = re.sub(r'[<>:"/\\|?*]+', "_", title).strip() or "试卷"
    return stem[:80]


def _mock_generate_exam_paper(arguments: dict[str, Any]) -> dict[str, Any]:
    title = arguments.get("title") or "试卷"
    questions = arguments.get("questions") or []
    doc = Document()
    doc.add_heading(title, 0)
    for i, q in enumerate(questions, 1):
        pts = q.get("points")
        head = f"{i}. {q.get('content', '')}"
        if pts is not None:
            head += f"  ({pts}分)"
        doc.add_paragraph(head)
        for j, opt in enumerate(q.get("options") or []):
            letter = chr(ord("A") + j)
            doc.add_paragraph(f"  {letter}. {opt}")
        if q.get("answer"):
            doc.add_paragraph(f"参考答案：{q['answer']}")
    path = _unique_paper_path(title, ".docx")
    doc.save(path)
    md_path = path.with_suffix(".md")
    md_path.write_text(_to_markdown(title, questions), encoding="utf-8")
    return {
        "docx_path": relpath(path),
        "md_path": relpath(md_path),
        "mock": True,
        "title": title,
    }


def _unique_paper_path(title: str, suffix: str) -> Path:
    out = file_root() / "papers"
    out.mkdir(parents=True, exist_ok=True)
    stem = _safe_stem(title)
    path = out / f"{stem}{suffix}"
    n = 1
    while path.exists():
        path = out / f"{stem}_{n}{suffix}"
        n += 1
    return path


def _to_markdown(title: str, questions: list[dict[str, Any]]) -> str:
    lines = [f"# {title}", ""]
    for i, q in enumerate(questions, 1):
        lines.append(f"## {i}")
        lines.append(
            f"kind: {'objective' if q.get('type') in {'single_choice','multiple_choice','true_false'} else 'subjective'}"
        )
        lines.append(f"score: {q.get('points', 0)}")
        lines.append(f"stem: {q.get('content', '')}")
        if q.get("options"):
            lines.append("options:")
            for j, opt in enumerate(q["options"]):
                lines.append(f"- {chr(ord('A')+j)}. {opt}")
        if q.get("answer"):
            lines.append(f"answer: {q['answer']}")
        lines.append("")
    return "\n".join(lines)


def _edupaper_command() -> list[str]:
    from zypg.config import settings

    cmd = (settings.edupaper_mcp_cmd or "").strip()
    if cmd:
        args = [a.strip() for a in (settings.edupaper_mcp_args or "").split(",") if a.strip()]
        return [cmd, *args]
    return default_npx_command("edupaper-mcp")


def _materialize_mcp_text(title: str, questions: list[dict[str, Any]], text: str) -> dict[str, Any]:
    from zypg.storage.files import is_placeholder_url

    url = extract_download_url(text)
    if not url or is_placeholder_url(url):
        raise RuntimeError(f"edupaper-mcp 未返回可用下载链接：{text[:500]}")
    path = _unique_paper_path(title, ".docx")
    _download_docx(url, path)
    md_path = path.with_suffix(".md")
    md_path.write_text(_to_markdown(title, questions), encoding="utf-8")
    return {
        "docx_path": relpath(path),
        "md_path": relpath(md_path),
        "remote_url": url,
        "mock": False,
        "title": title,
        "mcp_text": text,
    }


def _download_docx(url: str, dest: Path) -> None:
    with httpx.Client(timeout=60.0, follow_redirects=True) as client:
        resp = client.get(url, headers={"User-Agent": "zypg/0.1"})
        resp.raise_for_status()
        dest.write_bytes(resp.content)
    data = dest.read_bytes()
    if len(data) < 32 or data[:2] != b"PK":
        dest.unlink(missing_ok=True)
        raise RuntimeError(f"从 edupaper-mcp 下载的不是有效 DOCX：{url}")


def _result_text(result: dict[str, Any]) -> str:
    parts: list[str] = []
    for item in result.get("content") or []:
        if isinstance(item, dict) and item.get("type") == "text":
            parts.append(str(item.get("text") or ""))
    if parts:
        return "\n".join(parts).strip()
    if isinstance(result.get("raw"), str):
        return result["raw"]
    return json_dumps(result)


def json_dumps(value: Any) -> str:
    import json

    return json.dumps(value, ensure_ascii=False)


async def generate_exam_paper(arguments: dict[str, Any]) -> dict[str, Any]:
    from zypg.config import settings

    title = arguments.get("title") or "试卷"
    questions = normalize_exam_questions(arguments.get("questions") or [])
    payload = {"title": title, "questions": questions}
    if use_edupaper_mock():
        return _mock_generate_exam_paper(payload)

    url = (settings.edupaper_mcp_url or "").strip()
    timeout = float(settings.edupaper_mcp_timeout or 120)
    try:
        if url:
            text = await _call_http(url, payload, timeout)
        else:
            result = call_mcp_stdio(_edupaper_command(), "generate_exam_paper", payload, timeout=timeout)
            text = _result_text(result)
        return _materialize_mcp_text(title, questions, text)
    except Exception:
        out = _mock_generate_exam_paper(payload)
        out["mcp_fallback"] = True
        return out


async def _call_http(url: str, arguments: dict[str, Any], timeout: float) -> str:
    body = {
        "jsonrpc": "2.0",
        "id": 1,
        "method": "tools/call",
        "params": {"name": "generate_exam_paper", "arguments": arguments},
    }
    async with httpx.AsyncClient(timeout=timeout) as client:
        resp = await client.post(url, json=body)
        resp.raise_for_status()
        data = resp.json()
    if "error" in data:
        raise RuntimeError(data["error"])
    result = data.get("result") or data
    if not isinstance(result, dict):
        raise RuntimeError(f"edupaper-mcp HTTP 返回异常: {data}")
    if result.get("isError"):
        raise RuntimeError(_result_text(result) or "edupaper-mcp HTTP isError")
    return _result_text(result)


def parse_paper_file(path: str | Path) -> dict[str, Any]:
    path = Path(path)
    text = _read_paper_text(path)
    parsed = parse_markdown_paper(text)
    if _usable_items(parsed.get("items") or []):
        parsed.setdefault("name", path.stem)
        return parsed
    loose = _parse_loose_exam(text)
    if _usable_items(loose):
        title = _title_from_text(text) or path.stem
        return {"name": title, "subject": parsed.get("subject"), "items": loose}
    if path.suffix.lower() in {".docx", ".doc", ".pdf"} and (text or "").strip():
        return {
            "name": path.stem,
            "subject": None,
            "items": [
                {
                    "number": 1,
                    "kind": "subjective",
                    "stem": text[:4000],
                    "score": 10,
                    "knowledge_point": None,
                    "options": None,
                    "answer_key": None,
                    "rubric": ["请老师补评分点"],
                }
            ],
        }
    if path.suffix.lower() not in {".docx", ".doc", ".pdf"}:
        parsed.setdefault("name", path.stem)
        return parsed
    raise ValueError(f"无法从 {path.name} 解析出题目。请用带题号的 Word/Markdown，或先在试卷页核对格式。")


def _usable_items(items: list[dict[str, Any]]) -> bool:
    if len(items) >= 2:
        return True
    if len(items) == 1 and (items[0].get("stem") or "").strip() and items[0].get("options"):
        return True
    return bool(items) and all((it.get("stem") or "").strip() and it.get("kind") for it in items) and len(items) >= 1 and any(
        it.get("options") or it.get("answer_key") or it.get("rubric") for it in items
    )


def _title_from_text(text: str) -> str | None:
    for line in (text or "").splitlines()[:12]:
        s = line.strip().lstrip("#").strip()
        if 2 <= len(s) <= 40 and not re.match(r"^\d+", s):
            return s
    return None


def _read_paper_text(path: Path) -> str:
    suf = path.suffix.lower()
    if suf in {".docx", ".doc"}:
        return "\n".join(p.text for p in Document(str(path)).paragraphs)
    if suf == ".pdf":
        return _pdf_text(path)
    return path.read_text(encoding="utf-8-sig")


def _pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError("未安装 pypdf，无法解析 PDF。请改用 Word/Markdown，或 pip install pypdf。") from exc
    reader = PdfReader(str(path))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    if not text.strip():
        raise ValueError("PDF 没有可提取的文字（可能是扫描件）。请用 Word/Markdown，或先 OCR。")
    return text


def _parse_loose_exam(text: str) -> list[dict[str, Any]]:
    chunks = re.split(
        r"(?m)(?=^\s*(?:第\s*\d+\s*题|\d{1,3}\s*[\.、．]|[（(]\d{1,3}[）)]))",
        text or "",
    )
    items: list[dict[str, Any]] = []
    n = 0
    for chunk in chunks:
        body = chunk.strip()
        if len(body) < 6:
            continue
        if re.match(r"^(?:答案|参考答案|评分|注意事项)[:：]", body):
            continue
        n += 1
        opt_lines = re.findall(r"(?m)^\s*([A-Da-d])[\.、．\)]\s*(.+)$", body)
        options = None
        if len(opt_lines) >= 2:
            options = [f"{a.upper()}. {t.strip()}" for a, t in opt_lines]
        stem = body
        stem = re.sub(r"(?m)^\s*[A-Da-d][\.、．\)].+$", "", stem)
        stem = re.sub(r"(?s)(?:参考)?答案[:：].+$", "", stem)
        stem = re.sub(r"(?s)评分点[:：].+$", "", stem)
        stem = re.sub(r"^\s*(?:第\s*\d+\s*题|\d{1,3}\s*[\.、．]|[（(]\d{1,3}[）)])\s*", "", stem.strip())
        stem = re.sub(r"\n{3,}", "\n\n", stem).strip()
        ans_m = re.search(r"(?:参考)?答案[:：]\s*([A-Da-d]|[^\n]{1,80})", body)
        answer = (ans_m.group(1).strip() if ans_m else "") or ""
        rubric_m = re.findall(r"(?m)^\s*[-－]\s+(.+)$", body)
        kind = "objective" if options else "subjective"
        answer_key = None
        if answer:
            if kind == "objective":
                answer_key = {"letter": answer.upper()[:1]}
            else:
                answer_key = {"text": answer}
        if not stem:
            continue
        items.append(
            {
                "number": n,
                "kind": kind,
                "stem": stem[:2000],
                "score": 5 if kind == "objective" else 10,
                "knowledge_point": None,
                "options": options,
                "answer_key": answer_key,
                "rubric": rubric_m or (None if kind == "objective" else ["按步骤给分"]),
            }
        )
    return items


def parse_markdown_paper(text: str) -> dict[str, Any]:
    lines = text.replace("\r\n", "\n").split("\n")
    name = "未命名试卷"
    subject = None
    items: list[dict[str, Any]] = []
    cur: dict[str, Any] | None = None
    mode = None
    for line in lines:
        raw = line.rstrip()
        if raw.startswith("# "):
            name = raw[2:].strip()
            continue
        if raw.lower().startswith("subject:"):
            subject = raw.split(":", 1)[1].strip()
            continue
        if raw.startswith("## "):
            if cur:
                items.append(_finalize_item(cur, len(items) + 1))
            cur = {"number_hint": raw[3:].strip(), "options_list": [], "rubric_list": []}
            mode = None
            continue
        if cur is None:
            continue
        low = raw.strip()
        if low.startswith("kind:"):
            cur["kind"] = low.split(":", 1)[1].strip()
        elif low.startswith("knowledge_point:"):
            cur["knowledge_point"] = low.split(":", 1)[1].strip()
        elif low.startswith("score:"):
            try:
                cur["score"] = float(low.split(":", 1)[1].strip())
            except ValueError:
                cur["score"] = 0
        elif low.startswith("stem:"):
            cur["stem"] = low.split(":", 1)[1].strip()
            mode = "stem"
        elif low.startswith("answer:"):
            cur["answer"] = low.split(":", 1)[1].strip()
            mode = None
        elif low == "options:":
            mode = "options"
        elif low == "rubric:":
            mode = "rubric"
        elif low.startswith("- ") and mode == "options":
            cur["options_list"].append(low[2:].strip())
        elif low.startswith("- ") and mode == "rubric":
            cur["rubric_list"].append(low[2:].strip())
        elif raw.strip() and mode == "stem":
            cur["stem"] = (cur.get("stem") or "") + "\n" + raw.strip()
    if cur:
        items.append(_finalize_item(cur, len(items) + 1))
    return {"name": name, "subject": subject, "items": items}


def _finalize_item(cur: dict[str, Any], fallback_n: int) -> dict[str, Any]:
    kind = cur.get("kind") or ("objective" if cur.get("options_list") else "subjective")
    options = cur.get("options_list") or None
    answer = cur.get("answer")
    answer_key = None
    if answer is not None:
        if kind == "objective":
            letter = str(answer).strip().upper()[:1]
            answer_key = {"letter": letter}
        else:
            answer_key = {"text": answer}
    try:
        number = int(str(cur.get("number_hint", "")).split()[0])
    except (ValueError, IndexError):
        number = fallback_n
    return {
        "number": number,
        "kind": kind,
        "stem": cur.get("stem") or "",
        "score": cur.get("score") or 0,
        "knowledge_point": cur.get("knowledge_point"),
        "options": options,
        "answer_key": answer_key,
        "rubric": cur.get("rubric_list") or None,
    }
